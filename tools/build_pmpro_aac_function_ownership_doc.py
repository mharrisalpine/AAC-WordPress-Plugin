from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path('/Users/mharris/Desktop/WordPress Page Matched Web App/outputs/AAC_PMPro_Function_Ownership_Guide.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)

BLUE = '2E74B5'
DARK_BLUE = '1F4D78'
LIGHT_BLUE = 'E8EEF5'
GOLD = 'F8C235'
RED = 'B71C1C'
DARK = '16130F'
GRAY = '666666'
LIGHT_GRAY = 'F5F5F3'
WHITE = 'FFFFFF'

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

def set_run_font(run, size=None, bold=None, color=None, italic=None, name='Calibri'):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Calibri'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for style_name, size, color, before, after in [
    ('Heading 1', 16, BLUE, 18, 10),
    ('Heading 2', 13, BLUE, 14, 7),
    ('Heading 3', 12, DARK_BLUE, 10, 5),
]:
    st = styles[style_name]
    st.font.name = 'Calibri'
    st._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for list_name in ['List Bullet', 'List Number']:
    st = styles[list_name]
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    st.paragraph_format.left_indent = Inches(0.375)
    st.paragraph_format.first_line_indent = Inches(-0.188)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.25

if 'Callout' not in styles:
    callout = styles.add_style('Callout', WD_STYLE_TYPE.PARAGRAPH)
else:
    callout = styles['Callout']
callout.font.name = 'Calibri'
callout.font.size = Pt(11)
callout.font.bold = True
callout.font.color.rgb = RGBColor.from_string(DARK_BLUE)
callout.paragraph_format.left_indent = Inches(0.18)
callout.paragraph_format.right_indent = Inches(0.18)
callout.paragraph_format.space_before = Pt(8)
callout.paragraph_format.space_after = Pt(12)
callout.paragraph_format.line_spacing = 1.2

def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths_dxa)
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(total))
    tblW.set(qn('w:type'), 'dxa')
    tblInd = tblPr.find(qn('w:tblInd'))
    if tblInd is None:
        tblInd = OxmlElement('w:tblInd')
        tblPr.append(tblInd)
    tblInd.set(qn('w:w'), '120')
    tblInd.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:w'), str(widths_dxa[idx]))
            tcW.set(qn('w:type'), 'dxa')
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Page ')
    set_run_font(run, size=9, color=GRAY)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(fld)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = header.add_run('AMERICAN ALPINE CLUB  |  MEMBERSHIP PLATFORM REFERENCE')
set_run_font(hr, size=8.5, bold=True, color=GRAY)
add_page_number(section.footer.paragraphs[0])

def add_bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.add_run(text)
    return p

def add_number(text):
    p = doc.add_paragraph(style='List Number')
    p.add_run(text)
    return p

def add_bullets(items):
    for item in items:
        add_bullet(item)

def add_group(title, items):
    doc.add_heading(title, level=2)
    add_bullets(items)

def add_callout(label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_BLUE)
    p = cell.paragraphs[0]
    p.style = 'Callout'
    r = p.add_run(label + ': ')
    set_run_font(r, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_run_font(r2, bold=False, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(52)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('SYSTEM OWNERSHIP GUIDE')
set_run_font(r, size=10, bold=True, color=RED)
p.paragraph_format.space_after = Pt(16)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('PMPro and AAC Member App Functions')
set_run_font(r, size=28, bold=True, color=DARK_BLUE)
p.paragraph_format.space_after = Pt(8)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Native membership functions, AAC customizations, and shared ownership boundaries')
set_run_font(r, size=14, color=GRAY)
p.paragraph_format.space_after = Pt(28)

table = doc.add_table(rows=3, cols=2)
set_table_geometry(table, [2700, 6660])
metadata = [('Organization', 'American Alpine Club'), ('Platform', 'WordPress, Paid Memberships Pro, AAC Member Portal, Stripe, and Salesforce'), ('Purpose', 'Administration, development, migration planning, and troubleshooting')]
for i, (label, value) in enumerate(metadata):
    shade_cell(table.cell(i, 0), LIGHT_BLUE)
    table.cell(i, 0).text = label
    table.cell(i, 1).text = value
    for run in table.cell(i, 0).paragraphs[0].runs:
        set_run_font(run, bold=True, color=DARK_BLUE)
    for run in table.cell(i, 1).paragraphs[0].runs:
        set_run_font(run, color=DARK)

doc.add_paragraph().paragraph_format.space_after = Pt(18)
add_callout('Core rule', 'PMPro is the membership and payment system of record. The AAC Member App is the custom presentation, workflow, business-rules, family-management, reporting, and integration layer built around PMPro.')

doc.add_page_break()
doc.add_heading('How to Use This Guide', level=1)
doc.add_paragraph('This guide separates native PMPro behavior from custom AAC behavior. A function may be visually presented by the AAC application while still being calculated, validated, charged, or stored by PMPro.')
doc.add_heading('Ownership Labels', level=2)
add_bullets([
    'PMPro native: PMPro core supplies and owns the function.',
    'PMPro add-on: An installed PMPro extension supplies the function.',
    'AAC custom: The AAC Member Portal, React application, templates, services, or metadata supply the function.',
    'Hybrid: AAC controls the interface or business workflow, while PMPro remains authoritative for membership or financial records.',
    'External system: Stripe executes card transactions; Salesforce receives and returns permitted synchronized data.',
])
doc.add_heading('Contents', level=2)
for item in ['1. Native PMPro functions', '2. Custom AAC Member App functions', '3. Hybrid ownership matrix', '4. Administration and troubleshooting rules', '5. Reference sources']:
    add_bullet(item)

doc.add_page_break()
doc.add_heading('1. Native PMPro Functions', level=1)
doc.add_paragraph('PMPro core and the installed PMPro add-ons provide the underlying membership, checkout, payment, subscription, order, and access-control functions described below.')

native_groups = [
('Membership levels and prices', [
'Membership level records, names, descriptions, IDs, signup availability, and level groups.',
'Initial checkout payment, recurring billing amount, billing interval, billing-cycle limit, trials, and expiration duration.',
'Free, one-time, recurring, trial, fixed-term, and lifetime membership configurations.',
'Whether a member may hold one or multiple levels and the native rules for changing levels.',
'Administrative assignment, update, cancellation, upgrade, and downgrade of membership levels.',
'The selected PMPro level and checkout calculation are authoritative; marketing prices shown by React are not.',
]),
('Signup and account creation', [
'The native membership checkout form and direct checkout for a selected level ID.',
'Creation of the WordPress user during checkout.',
'Username, email, password, duplicate-account, and required-field validation.',
'Recognition of a logged-in member completing checkout.',
'Assignment of the purchased membership, membership start date, and membership-history entry.',
'Checkout-completion hooks, welcome messages, receipts, and confirmation routing.',
]),
('User and checkout fields', [
'Registration of user fields and field groups.',
'Required and optional field status.',
'Supported text, select, checkbox, radio, date, and related field types.',
'Server-side validation, sanitization, and storage in WordPress user metadata.',
'Checkout, profile, and administrative display locations.',
]),
('Billing address', [
'Billing first and last name, address lines, country, city, state or province, postal code, email, and phone where enabled.',
'Saving billing information with the order and passing it to the configured payment gateway.',
]),
('Checkout calculation and payment processing', [
'Loading the selected level and calculating its initial and recurring charges.',
'Applying valid PMPro discount codes and supported price-adjusting add-ons.',
'Building the authoritative checkout total.',
'Connecting checkout to Stripe and handling success or failure responses.',
'Creating the PMPro order and storing payment and subscription transaction identifiers.',
'Recording gateway, environment, payment status, card type, and last four digits when supplied by the gateway.',
'Associating the order with the WordPress user and membership.',
]),
('Discount codes', [
'Code value, availability dates, usage limits, eligible levels, and validation.',
'Discounted initial payment, recurring amount, trial terms, and expiration terms.',
'Association of the accepted code with the membership and PMPro order.',
]),
('Recurring subscriptions and automatic renewal', [
'Creation and tracking of the recurring Stripe subscription.',
'Billing amount, interval, subscription transaction ID, and recurring-payment orders.',
'Processing gateway webhooks and maintaining access after successful renewal payments.',
'Cancellation of supported gateway subscriptions and reflection of subscription status in PMPro.',
'Native billing-information update and renewal communications.',
]),
('Manual renewal and expiration', [
'Fixed-term start and end dates, expiration state, and the end of protected-content access.',
'Renew links, renewal through checkout, expiration warnings, and renewal emails.',
'Membership-history states for renewed, changed, cancelled, expired, and inactive records.',
]),
('Cancellation', [
'Member-requested and administrative cancellation.',
'Gateway-subscription cancellation where supported.',
'Membership-status and history updates, access ending, and cancellation communications.',
]),
('Upgrades, downgrades, and proration', [
'Changing levels through checkout and replacing a previous level in one-per-group configurations.',
'Assignment of the new level and creation of the associated order and membership-history record.',
'Proration supplied by the installed PMPro add-on or PMPro server-side hooks.',
]),
('Orders, subscriptions, and membership history', [
'Current and historical membership records, start and end dates, status, and copied term pricing.',
'Orders, subscriptions, discount-code association, transaction IDs, gateway, billing address, payment status, and checkout date.',
'Separation of membership access records, payment orders, and gateway subscription records.',
]),
('Confirmation, receipts, and email', [
'Checkout confirmation, order receipt data, welcome messages, payment receipts, recurring-payment messages, expiration warnings, and cancellation messages.',
]),
('Content access', [
'Restriction of posts, pages, categories, and other content by membership level.',
'Level checks, protected-content messages, redirects, and expiration of access.',
]),
('PMPro administration', [
'Members, levels, orders, subscriptions, discount codes, reports, payment settings, pages, email templates, user fields, and design settings.',
'Manual membership changes, manual orders, gateway configuration, and live/test environments.',
]),
('Group Accounts add-on', [
'Parent and child group memberships, group levels, seat capacity, and linking children to a paying parent.',
]),
('Donations add-on', [
'Donation value entering checkout, adjustment of the PMPro order total, gateway charge inclusion, and order storage.',
]),
('Renewal and proration add-ons', [
'Upgrade proration, immediate paid level changes, renewal controls, and supported scheduled changes where installed and configured.',
]),
]
for title, items in native_groups:
    add_group(title, items)

doc.add_page_break()
doc.add_heading('2. Custom AAC Member App Functions', level=1)
doc.add_paragraph('The AAC Member App customizes the member experience and adds AAC-specific workflows, data normalization, administration, and integrations. It does not replace PMPro as the final membership and financial authority.')

custom_groups = [
('Custom signup presentation', [
'Branded Select Your Plan cards, AAC typography, colors, outlines, images, benefit lists, selected states, and card ordering.',
'Mapping a selected membership card to the correct PMPro level ID.',
'Four-stage signup experience: Account Information, Member Information, Publication Preferences, and Discounts/Promo/Checkout.',
'Custom Continue and Back navigation, responsive layout, header/footer integration, and checkout height/scroll behavior.',
]),
('AAC tier and membership-card display', [
'Supporter, Partner, Leader, Advocate, and Partner Family presentation.',
'AAC-specific level-label normalization, benefit copy, status badges, plan-type badges, family badges, membership number, Member Since, Valid Through, and auto-renewal display.',
'Topographic visual membership card and responsive profile presentation.',
]),
('Eligibility and conditional checkout', [
'Skipping Publication Preferences for Supporter and hiding print/T-shirt options for ineligible international members.',
'Conditional stages and fields based on membership level, country, discount selection, and family selection.',
'International fulfillment notices and country-routed AAC behavior.',
]),
('Publication preferences', [
'AAC cards and images for the American Alpine Journal, Accidents in North American Climbing, American Climbing Journal, and Guidebook to Membership.',
'Click-to-select Print behavior, Digital default, card borders, and replacement of native dropdown presentation.',
'Eligibility by membership level and country, profile messaging, download buttons, and mapping into saved preference fields.',
]),
('T-shirt preference behavior', [
'Placement beside Phone, conditional visibility, AAC option labels, No T-shirt handling for ineligible members, and prevention of duplicate controls.',
'Synchronization of the configured PMPro field to the AAC canonical profile value.',
]),
('AAC membership-discount cards', [
'Student, Military, Family, and No Discount card presentation.',
'Conditional school/university, service-component, and family-detail inputs.',
'Mapping Student and Military selections to configured PMPro discount codes.',
]),
('Promo-code presentation', [
'Placement below Make a Gift and above Order Summary, stable placement after other selections, side-by-side layout with Automatic Renewal, and AAC styling.',
]),
('Make a Gift interface', [
'No thanks, preset amounts, custom amount, selected-button styling, placement, and separate donation line-item presentation.',
]),
('Custom order summary', [
'Itemized membership, discount, family addition, publication addition, donation, proration, renewal, and total presentation based on PMPro checkout state.',
]),
('Upgrade and renewal workflow', [
'Upgrade/Renew buttons, upgrade-level cards, current-level display, locked downgrades, selected-level routing, and shortened logged-in checkout.',
'Prepopulated member data, hidden account-creation fields for logged-in members, consolidated details/payment/summary, and custom intent flags.',
'Displaying PMPro-provided proration and returning members to the appropriate profile route.',
]),
('AAC early-renewal extensions', [
'Allowing eligible non-auto-renewing members to renew before expiration.',
'Intent to extend the existing expiration by another year while preserving the prior month/day.',
'Displaying the resulting Valid Through date and avoiding false automatic-renewal presentation.',
]),
('Scheduled downgrade behavior', [
'Eligibility rules, locked/selectable cards, renewal-boundary scheduling, pending-downgrade display, and cancellation/change of pending requests where supported.',
]),
('Family membership customization', [
'Redeem Family Membership Invite, invite codes, linked-member screens, child-account creation, email checks, seat presentation, dependent options, and AAC capacity defaults.',
'AAC parent/child metadata, imported relationship repair, administrative family tools, family badges, and checkout adjustments.',
]),
('Member Profile application', [
'Member Profile, Settings, Billing, Benefits, Contact Us, Sign Out, full-width navigation, active underline, header offsets, and responsive layout.',
'Profile editing, publication settings, benefits, linked accounts, membership management, upgrade/renew/cancel/billing actions, transaction register, and custom receipt display.',
]),
('Billing integration', [
'Embedding PMPro billing updates in the AAC shell, suppressing duplicate native navigation, combining member/payment sections, matching AAC styling, and dynamic embed sizing.',
]),
('Transaction register', [
'Branded PMPro order display with individual transaction boxes, dues/donation separation, discounts, totals, dates, and order/gateway references.',
]),
('Custom profile data and REST API', [
'Normalization of address, phone, T-shirt, publications, membership card, member-since, family, benefit, and Salesforce-related data.',
'Custom authenticated endpoints for profile, transactions, membership actions, downgrade, family operations, member database, settings, publications, and benefits.',
'Nonce injection, authenticated requests, session-aware loading, login/reset integration, and cache invalidation.',
]),
('Member database and administration', [
'Combined PMPro/profile/family views, CSV export, import extensions, membership history and order import, relationship repair, diagnostic tools, and data dictionaries.',
'Daily member-auth export over SFTP and administrative operational tools.',
]),
('Benefits system', [
'Level-based benefit cards, gallery content, partner links, ExpertVoice presentation, configurable images/copy, and server-side eligibility normalization.',
]),
('Confirmation and redirects', [
'Post-purchase routing, sign-in redirect where required, return to the member profile after logged-in changes, branded confirmation, and custom receipt/donation display.',
]),
('Caching and release behavior', [
'Normalized profile caching and invalidation after profile, PMPro, family, or settings changes.',
'Protection from cached authenticated responses, checkout embed sizing, runtime WordPress configuration, and frontend cache-busting.',
]),
('Salesforce integration - separate AAC plugin', [
'PMPro/AAC event listeners, outbound queue, Salesforce OAuth and REST upserts.',
'Contact, membership-term, membership-dues, and donation-line synchronization.',
'External IDs, idempotency, retries, dead-letter processing, field-level error logs, reconciliation, and health monitoring.',
'Optional permitted Salesforce-to-WordPress profile and membership updates.',
'Rejection of inbound changes to PMPro payments, Stripe identifiers, orders, donations, and discounts.',
]),
]
for title, items in custom_groups:
    add_group(title, items)

doc.add_page_break()
doc.add_heading('3. Hybrid Ownership Matrix', level=1)
doc.add_paragraph('The following functions cross the AAC and PMPro boundary. The AAC side controls presentation and workflow; the PMPro side remains authoritative for membership and financial processing.')
matrix = [
('Plan selection', 'Cards and navigation', 'Level ID, availability, and checkout level'),
('Price display', 'Marketing/card presentation', 'Authoritative checkout price'),
('Signup fields', 'Layout and conditional visibility', 'Registration, validation, and persistence'),
('Address fields', 'Two-column branded layout', 'Billing data and gateway submission'),
('T-shirt', 'Placement and synchronization', 'Registered required-field validation'),
('Publications', 'Card interface and eligibility', 'Registered field storage'),
('Student/Military', 'Selection cards and supporting fields', 'Discount-code validation and financial application'),
('Promo code', 'Placement and styling', 'Validation and price override'),
('Donation', 'Buttons and summary line', 'Donation amount and order total'),
('Order summary', 'Itemized presentation', 'Final calculated amount and order'),
('Card payment', 'Styled checkout container', 'Gateway request and order creation'),
('Stripe charge', 'No direct charge execution', 'PMPro coordinates; Stripe executes'),
('Auto-renewal', 'Toggle presentation', 'Subscription creation and recurring processing'),
('Upgrade', 'Cards and routing', 'Level change, order, and subscription'),
('Proration', 'Display of returned amount', 'PMPro add-on or hook calculation'),
('Manual renewal', 'Custom entry and renewal intent', 'Checkout, order, and stored term'),
('Expiration', 'Membership-card display', 'Authoritative end date and status'),
('Cancellation', 'Branded route and presentation', 'Subscription and membership cancellation'),
('Family account', 'AAC invite and linked-account interface', 'Group Accounts membership records'),
('Billing update', 'Embedded branded page', 'Payment-method update'),
('Transaction register', 'Branded breakdown', 'PMPro order records'),
('Salesforce', 'Queue, transformation, and field authority', 'Supplies membership/order source events'),
]
table = doc.add_table(rows=1, cols=3)
set_table_geometry(table, [1900, 3550, 3910])
for idx, text in enumerate(['Function', 'AAC Member App owns', 'PMPro owns']):
    cell = table.rows[0].cells[idx]
    cell.text = text
    shade_cell(cell, LIGHT_BLUE)
    for run in cell.paragraphs[0].runs:
        set_run_font(run, bold=True, color=DARK_BLUE)
table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
for i, row in enumerate(matrix):
    cells = table.add_row().cells
    for j, text in enumerate(row):
        cells[j].text = text
        if i % 2 == 1:
            shade_cell(cells[j], LIGHT_GRAY)
        for run in cells[j].paragraphs[0].runs:
            set_run_font(run, size=9.5, bold=(j == 0), color=DARK)
set_table_geometry(table, [1900, 3550, 3910])

doc.add_page_break()
doc.add_heading('4. Administration and Troubleshooting Rules', level=1)
rules = [
'Configure levels, authoritative prices, recurrence, trials, expiration, discount codes, gateways, subscriptions, and PMPro-required fields in PMPro.',
'Configure AAC card copy, imagery, benefit text, layout, conditional presentation, and member-profile experience in the AAC Member Portal.',
'Never repair a financial discrepancy with display-only JavaScript. Reconcile the PMPro level, order, gateway transaction, discount, donation, and subscription.',
'All required PMPro controls must remain inside the submitted PMPro form with their original validation and server mapping intact.',
'Treat PMPro orders and gateway transaction IDs as authoritative when totals differ from the visual order summary.',
'Inspect both PMPro Group Accounts and AAC relationship metadata before repairing family members.',
'Keep direct Salesforce calls out of React. The separate Salesforce Sync plugin owns CRM communication.',
'Keep PMPro payment, order, donation, discount, Stripe, and subscription data outbound-only to Salesforce.',
'A local React build cannot prove WordPress authentication, PMPro checkout, Stripe, Group Accounts, or Salesforce behavior. Verify on the target WordPress environment.',
'A packaged ZIP or Git push is not a deployment. Verify the active plugin version and the actual production workflow.',
]
for rule in rules:
    add_bullet(rule)

doc.add_heading('Decision Guide: Where Should a Change Be Made?', level=2)
decisions = [
('Change a membership price or billing interval', 'PMPro level settings'),
('Change card copy, benefit text, or card styling', 'AAC Member Portal settings/code'),
('Add a new required checkout field', 'PMPro registration plus AAC layout/storage mapping'),
('Change a promo-code amount or eligibility', 'PMPro discount-code settings'),
('Change Student/Military card appearance', 'AAC Member Portal'),
('Change the final amount charged', 'PMPro/add-on/server-side calculation; never React alone'),
('Change family seat capacity', 'PMPro Group Accounts plus AAC family configuration'),
('Change Salesforce field authority or mapping', 'AAC Salesforce Sync field maps'),
('Change payment method or gateway', 'PMPro payment settings and Stripe connection'),
('Change renewal or expiration rules', 'PMPro level/add-on configuration and any AAC renewal hook'),
]
table = doc.add_table(rows=1, cols=2)
set_table_geometry(table, [3600, 5760])
for idx, text in enumerate(['Requested change', 'Primary owner/location']):
    table.cell(0, idx).text = text
    shade_cell(table.cell(0, idx), LIGHT_BLUE)
    for run in table.cell(0, idx).paragraphs[0].runs:
        set_run_font(run, bold=True, color=DARK_BLUE)
for i, row in enumerate(decisions):
    cells = table.add_row().cells
    cells[0].text, cells[1].text = row
    if i % 2 == 1:
        shade_cell(cells[0], LIGHT_GRAY); shade_cell(cells[1], LIGHT_GRAY)
set_table_geometry(table, [3600, 5760])

doc.add_page_break()
doc.add_heading('5. Reference Sources', level=1)
sources = [
('AAC Member App Administrator, PMPro, and Developer Guide', 'docs/AAC_MEMBER_APP_OPERATIONS_GUIDE.md'),
('AAC Salesforce Sync', 'wordpress/aac-salesforce-sync/README.md'),
('AAC Salesforce Sync Production Runbook', 'wordpress/aac-salesforce-sync/PRODUCTION-RUNBOOK.md'),
('PMPro Membership Levels', 'https://www.paidmembershipspro.com/documentation/membership-levels/'),
('PMPro Membership Checkout', 'https://www.paidmembershipspro.com/documentation/frontend-pages/membership-checkout/'),
('PMPro Renewal Guide', 'https://www.paidmembershipspro.com/renewal-guide/'),
('PMPro Database Structure', 'https://www.paidmembershipspro.com/documentation/advanced/pmpro-database-structure/'),
('PMPro Discount Codes', 'https://www.paidmembershipspro.com/documentation/settings/discount-codes/'),
('PMPro Group Accounts', 'https://www.paidmembershipspro.com/add-ons/group-accounts'),
]
for name, location in sources:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(name + ': ')
    set_run_font(r, bold=True, color=DARK_BLUE)
    r = p.add_run(location)
    set_run_font(r, color=DARK)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('End of guide')
set_run_font(r, size=9, italic=True, color=GRAY)

# Keep headings and table rows readable across page breaks.
for p in doc.paragraphs:
    if p.style and p.style.name.startswith('Heading'):
        p.paragraph_format.keep_with_next = True
for table in doc.tables:
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement('w:cantSplit')
        trPr.append(cant_split)

doc.core_properties.title = 'PMPro and AAC Member App Functions'
doc.core_properties.subject = 'Native PMPro and custom AAC Member App function ownership'
doc.core_properties.author = 'American Alpine Club'
doc.core_properties.keywords = 'PMPro, AAC Member App, WordPress, membership, Salesforce, Stripe'
doc.save(OUT)
print(OUT)
